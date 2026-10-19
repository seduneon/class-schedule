"""Legacy converter: the previous capstone's wide Excel files -> SPEC §8 tables (SPEC §8.1).

Inputs (one folder): ``courses.xlsx`` (one row per section), ``professors.xlsx``,
``time_room.xlsx`` (one row per usable days/time/room) and ``preferences.xlsx`` (the
Qualtrics export: three header rows, then one row per response). Everything that needs
a human decision goes to ``review.csv``; nothing in free text is interpreted.
"""

from __future__ import annotations

import csv
import hashlib
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

from scheduler.domain import SectionType
from scheduler.io.schemas import DEFAULT_UNITS, day_family

Record = dict[str, object]

_SLOT = re.compile(
    r"^(?P<days>[MTWRFSU ]+?)\s*(?P<start>\d{1,2}:\d{2}\s*[AP]M)\s*-\s*(?P<end>\d{1,2}:\d{2}\s*[AP]M)$",
    re.IGNORECASE,
)
_TIME = re.compile(r"^(?P<start>\d{1,2}:\d{2}\s*[AP]M)\s*-\s*(?P<end>\d{1,2}:\d{2}\s*[AP]M)$", re.I)
_RECITATION = re.compile(r"recitation|reciation", re.IGNORECASE)

MULTI_SECTION = {1: 1, 2: -1, 3: 0}
"""Survey Q2 choice -> multi_section_pref (Yes / No / No preference)."""
DAY_CHOICE = {
    1: ("MWF", "allow"),
    2: ("TR", "allow"),
    3: ("MWF;TR", "avoid"),
    4: ("MWF;TR", "allow"),
}
"""Survey Q3 choice -> (day_families, day_family_mixing)."""
BREAK_CHOICE = {1: "spread", 2: "none"}
"""Survey Q6 choice -> schedule_style."""


@dataclass(frozen=True)
class ExtraOffering:
    """``--add-offering COURSE:TYPE:N:CAP``: sections missing from ``courses.xlsx``."""

    course_id: str
    section_type: SectionType
    num_sections: int
    capacity: int

    @classmethod
    def parse(cls, text: str) -> ExtraOffering:
        parts = [p.strip() for p in text.split(":")]
        if len(parts) != 4:
            raise ValueError(f"expected COURSE:TYPE:N:CAP, got {text!r}")
        course, kind, n, cap = parts
        return cls(course_id(course), SectionType(kind.upper()), int(n), int(cap))


@dataclass
class LegacyResult:
    """Converted tables (SPEC §8) plus the rows a human must review."""

    tables: dict[str, list[Record]]
    review: list[Record] = field(default_factory=list)


def course_id(code: str) -> str:
    """``MATH 571 (MATH 721) `` -> ``MATH571`` (cross-listing suffix dropped)."""
    return code.split("(")[0].strip().replace(" ", "")


def clock_24(text: str) -> str:
    """``01:30 PM`` -> ``13:30``."""
    t = datetime.strptime(re.sub(r"\s+", " ", text.strip()).upper(), "%I:%M %p")
    return t.strftime("%H:%M")


def pattern_id(family: str, start: str) -> str:
    """Pattern ID from family and 24-hour start, e.g. ``MWF0900``."""
    return f"{family}{start.replace(':', '')}"


def _sheet(path: Path) -> list[tuple[Any, ...]]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        return [tuple(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def _str(v: object) -> str:
    return "" if v is None else str(v).strip()


def _num(v: object) -> float | None:
    text = _str(v)
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _header_rows(rows: list[tuple[Any, ...]]) -> tuple[tuple[Any, ...], list[dict[str, Any]]]:
    """Header and non-empty records (first row = header)."""
    header = rows[0]
    out = []
    for r in rows[1:]:
        if any(_str(v) for v in r):
            out.append({_str(h): v for h, v in zip(header, r, strict=False) if _str(h)})
    return header, out


# ----- courses, rooms, patterns ------------------------------------------------------------


def _offerings(rows: list[dict[str, Any]], extra: tuple[ExtraOffering, ...]) -> list[Record]:
    counts: Counter[tuple[str, str, int]] = Counter()
    titles: dict[str, str] = {}
    for r in rows:
        cid = course_id(_str(r["Course_code"]))
        kind = _str(r["Section_type"]).upper()
        counts[(cid, kind, int(_num(r["Capacity"]) or 0))] += 1
        titles.setdefault(cid, _str(r["Course_title"]))
    for e in extra:
        counts[(e.course_id, e.section_type.value, e.capacity)] += e.num_sections
    out: list[Record] = []
    for (cid, kind, cap), n in sorted(counts.items()):
        digits = "".join(ch for ch in cid if ch.isdigit())
        out.append(
            {
                "course_id": cid,
                "title": titles.get(cid, cid),
                "level": int(digits[0]) * 100 if digits else 0,
                "section_type": kind,
                "num_sections": n,
                "capacity": cap,
                "units": DEFAULT_UNITS[SectionType(kind)],
                "room_type": "classroom",
            }
        )
    return out


@dataclass(frozen=True)
class _Slot:
    family: str
    start: str
    end: str

    @property
    def id(self) -> str:
        return pattern_id(self.family, self.start)


def _slot(days: str, time_range: str) -> _Slot:
    m = _TIME.match(time_range.strip())
    if m is None:
        raise ValueError(f"cannot parse time range {time_range!r}")
    return _Slot(day_family(days.replace(" ", "")), clock_24(m["start"]), clock_24(m["end"]))


def _time_room(rows: list[dict[str, Any]], review: list[Record]) -> dict[str, list[Record]]:
    slots: dict[str, _Slot] = {}
    rooms: dict[str, int] = {}
    pairs: set[tuple[str, str]] = set()
    for r in rows:
        if not _str(r.get("room")):
            continue
        slot = _slot(_str(r["day"]), _str(r["time"]))
        slots[slot.id] = slot
        rid = _str(r["room"])
        cap = int(_num(r["cap"]) or 0)
        if rid in rooms and rooms[rid] != cap:
            review.append(
                _rev("room_capacity", "", f"room {rid}: capacities {rooms[rid]} and {cap}")
            )
        rooms[rid] = max(cap, rooms.get(rid, 0))
        pairs.add((rid, slot.id))
    return {
        "time_patterns": [
            {
                "pattern_id": s.id,
                "days": s.family,
                "start": s.start,
                "end": s.end,
                "allowed_section_types": "L;R",
            }
            for s in sorted(slots.values(), key=lambda s: s.id)
        ],
        "rooms": [
            {
                "room_id": rid,
                "building": rid.split(".")[0] if "." in rid else rid,
                "capacity": cap,
                "room_type": "classroom",
            }
            for rid, cap in sorted(rooms.items())
        ],
        "room_availability": [{"room_id": r, "pattern_id": t} for r, t in sorted(pairs)],
    }


# ----- professors --------------------------------------------------------------------------


def _rev(kind: str, professor_id: str, detail: str) -> Record:
    return {"kind": kind, "professor_id": professor_id, "detail": detail}


@dataclass(frozen=True)
class _Prof:
    id: str
    first: str
    last: str
    email: str


def _professor_id(first: str, last: str, email: str) -> str:
    if email:
        return email.split("@")[0].lower()
    return re.sub(r"[^a-z0-9]+", ".", f"{first} {last}".lower()).strip(".")


def _professors(
    rows: list[dict[str, Any]], review: list[Record]
) -> tuple[list[Record], list[Record], list[_Prof]]:
    profs: list[Record] = []
    pre: list[Record] = []
    who: list[_Prof] = []
    for r in rows:
        first, last, email = _str(r["first_name"]), _str(r["second_name"]), _str(r.get("email"))
        if not (first or last):
            continue
        pid = _professor_id(first, last, email)
        load = _num(r.get("load")) or 0.0
        cpts, tpts = _num(r.get("course_points")), _num(r.get("time_points"))
        if cpts is not None and tpts is not None and cpts != tpts:
            review.append(
                _rev(
                    "priority",
                    pid,
                    f"course_points {cpts:g} != time_points {tpts:g}; used course_points",
                )
            )
        profs.append(
            {
                "professor_id": pid,
                "name": f"{first} {last}".strip(),
                "load_min": load,
                "load_max": load,
                "priority": (cpts or 100.0) / 100.0,
                "day_family_mixing": "forbid" if _num(r.get("MWForTR")) == 1 else "allow",
            }
        )
        who.append(_Prof(pid, first, last, email.lower()))
        for i in range(1, 10):
            course = _str(r.get(f"course{i}"))
            if not course:
                continue
            section = _str(r.get(f"section{i}"))
            label = f"{int(float(section))}L" if _num(section) is not None else ""
            pre.append(
                {
                    "professor_id": pid,
                    "course_id": course_id(course),
                    "section_type": "L",
                    "count": 1,
                    "section_label": label,
                }
            )
    return profs, pre, who


# ----- survey ------------------------------------------------------------------------------


@dataclass(frozen=True)
class _Column:
    index: int
    question: str  # Q1, Q2, ...
    course: str = ""
    recitation: bool = False
    slot: _Slot | None = None


def _survey_columns(codes: tuple[Any, ...], texts: tuple[Any, ...]) -> list[_Column]:
    out = []
    for i, code in enumerate(codes):
        q = _str(code)
        if not re.match(r"^Q\d+(_\d+)?$", q):
            continue
        text = _str(texts[i] if i < len(texts) else "")
        tail = text.rsplit(" - ", 1)[-1].strip()
        question = q.split("_")[0]
        if question == "Q1":
            course = tail.split("\t")[0].strip()
            out.append(_Column(i, question, course_id(course), bool(_RECITATION.search(tail))))
        elif question in ("Q4", "Q5"):
            m = _SLOT.match(re.sub(r"\s+", " ", tail))
            slot = _slot(m["days"], f"{m['start']}-{m['end']}") if m else None
            out.append(_Column(i, question, slot=slot))
        else:
            out.append(_Column(i, question))
    return out


def _match(resp: dict[str, str], who: list[_Prof]) -> tuple[_Prof | None, str]:
    email = resp["email"].lower()
    for p in who:
        if email and p.email == email:
            return p, ""
    same_last = [p for p in who if p.last.lower() == resp["last"].lower()]
    if len(same_last) == 1:
        return same_last[0], "survey response matched by last name (no e-mail match)"
    return None, ""


def _survey(
    rows: list[tuple[Any, ...]], who: list[_Prof], review: list[Record]
) -> tuple[dict[str, Record], list[Record], list[Record]]:
    """Per professor: professor-table overrides, course and time preference rows."""
    codes, texts = rows[1], rows[2]
    name_at = {_str(c): i for i, c in enumerate(codes)}
    cols = _survey_columns(codes, texts)
    rec_courses = {c.course for c in cols if c.question == "Q1" and c.recitation}

    latest: dict[str, tuple[datetime, tuple[Any, ...]]] = {}
    for row in rows[3:]:
        rid = _str(row[name_at["ResponseId"]]) if "ResponseId" in name_at else ""
        if not rid.startswith("R_"):
            continue
        resp = {
            "email": _str(row[name_at["RecipientEmail"]]),
            "last": _str(row[name_at["RecipientLastName"]]),
        }
        prof, note = _match(resp, who)
        if prof is None:
            review.append(_rev("unmatched_response", "", f"response {rid} matches no professor"))
            continue
        if note:
            review.append(_rev("name_match", prof.id, note))
        when = row[name_at["RecordedDate"]]
        stamp = when if isinstance(when, datetime) else datetime.min
        if prof.id in latest:
            review.append(
                _rev("duplicate_response", prof.id, "answered twice; latest response used")
            )
            if stamp <= latest[prof.id][0]:
                continue
        latest[prof.id] = (stamp, row)

    overrides: dict[str, Record] = {}
    cprefs: list[Record] = []
    tprefs: list[Record] = []
    for pid in sorted(latest):
        row = latest[pid][1]
        over: Record = {}
        for col in cols:
            value = _num(row[col.index])
            if col.question == "Q1" and value is not None:
                kind = "R" if col.recitation else ("L" if col.course in rec_courses else "")
                cprefs.append(
                    {
                        "professor_id": pid,
                        "course_id": col.course,
                        "section_type": kind,
                        "rank": int(value),
                    }
                )
            elif col.question in ("Q4", "Q5") and value is not None and col.slot is not None:
                tprefs.append({"professor_id": pid, "pattern_id": col.slot.id, "rank": int(value)})
            elif col.question == "Q2" and value is not None:
                over["multi_section_pref"] = MULTI_SECTION.get(int(value), 0)
            elif col.question == "Q3" and value is not None:
                fams, mix = DAY_CHOICE.get(int(value), ("", "allow"))
                over["day_families"], over["day_family_mixing"] = fams, mix
            elif col.question == "Q6" and value is not None:
                over["schedule_style"] = BREAK_CHOICE.get(int(value), "none")
            elif col.question == "Q7" and _str(row[col.index]):
                review.append(_rev("comment", pid, _str(row[col.index])))
        overrides[pid] = over
    return overrides, cprefs, tprefs


# ----- entry points ------------------------------------------------------------------------


def convert_legacy(
    folder: Path | str, extra: tuple[ExtraOffering, ...] = (), anonymize: bool = False
) -> LegacyResult:
    """Convert a legacy folder (SPEC §8.1); ``anonymize`` renames professors ``P01, ...``."""
    folder = Path(folder)
    review: list[Record] = []
    _h, course_rows = _header_rows(_sheet(folder / "courses.xlsx"))
    _h, room_rows = _header_rows(_sheet(folder / "time_room.xlsx"))
    _h, prof_rows = _header_rows(_sheet(folder / "professors.xlsx"))
    profs, pre, who = _professors(prof_rows, review)
    overrides, cprefs, tprefs = _survey(_sheet(folder / "preferences.xlsx"), who, review)
    for p in profs:
        pid = str(p["professor_id"])
        over = overrides.get(pid)
        if over is None:
            review.append(_rev("no_response", pid, "no survey response; no preferences"))
            continue
        forbid = p["day_family_mixing"] == "forbid"
        p.update(over)
        if forbid:
            p["day_family_mixing"] = "forbid"
    tables: dict[str, list[Record]] = {
        "professors": profs,
        "offerings": _offerings(course_rows, extra),
        **_time_room(room_rows, review),
        "course_preferences": cprefs,
        "time_preferences": tprefs,
        "pre_assignments": pre,
    }
    result = LegacyResult(tables, review)
    return _anonymized(result) if anonymize else result


def _anonymized(result: LegacyResult) -> LegacyResult:
    """Replace professor IDs and names; drop e-mail-derived IDs and free-text comments."""
    # numbered in hash order, not alphabetically, so P01.. does not follow the staff list
    ids = sorted(
        (str(p["professor_id"]) for p in result.tables["professors"]),
        key=lambda pid: hashlib.sha256(pid.encode()).hexdigest(),
    )
    width = max(2, len(str(len(ids))))
    new = {pid: f"P{i + 1:0{width}d}" for i, pid in enumerate(ids)}
    tables: dict[str, list[Record]] = {}
    for name, records in result.tables.items():
        out = []
        for rec in records:
            rec = dict(rec)
            if "professor_id" in rec:
                rec["professor_id"] = new[str(rec["professor_id"])]
            if name == "professors":
                rec["name"] = f"Professor {str(rec['professor_id'])[1:]}"
            out.append(rec)
        if name == "professors":
            out.sort(key=lambda r: str(r["professor_id"]))
        tables[name] = out
    review = [
        {**r, "professor_id": new.get(str(r["professor_id"]), "")}
        for r in result.review
        if r["kind"] not in ("comment", "unmatched_response")
    ]
    return LegacyResult(tables, review)


def write_legacy(result: LegacyResult, out: Path | str) -> Path:
    """Write the converted tables as CSV files plus ``review.csv``."""
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    for name, records in result.tables.items():
        if records:
            _write_csv(out / f"{name}.csv", records)
    _write_csv(out / "review.csv", result.review, ("kind", "professor_id", "detail"))
    return out


def _write_csv(path: Path, records: list[Record], columns: tuple[str, ...] = ()) -> None:
    cols = list(columns)
    for rec in records:
        cols += [k for k in rec if k not in cols]
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=cols)
        writer.writeheader()
        for rec in records:
            writer.writerow({k: rec.get(k, "") for k in cols})
